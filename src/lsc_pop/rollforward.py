"""Stage D - Roll forward the 2021 base to the reference mid-year (ADR-0004, ADR-0017).

1. Read the mid-year LSOA × sex × single-year estimates (S5; ``config.mye``) and check them
   (ROL-01 to ROL-04, including the accredited broad-age file S5b).
2. Turn the 2021 base into **ethnic shares** for each LSOA × sex × target age, using
   ``source_ages`` below. Pooled 2021 counts are summed over the source ages, then divided by
   their total across ethnic groups.

   * ``cohort`` (default): target age a takes the 2021 cohort now aged a, i.e. 2021 age a − k with
     k = ``config.shift_years`` (3 for mid-2024). Target 90+ pools 2021 ages 90 − k ... 90.
     Targets 0 ... k−1 (born after Census day) use ``rollforward.newborn_proxy_ages``.
   * ``static``: target age a takes 2021 age a (sensitivity variant).

3. Where the 2021 LSOA × sex has nobody at the source ages, the shares fall back, in order, to:
   (1) the LSOA × sex pooled over the RM032 band of the source ages; (2) the LSOA × sex, all ages;
   (3) the LSOA, both sexes, all ages; (4) the LTLA × sex at the source ages. The level used is
   recorded per cell and summarised (ROL-08).
4. Estimates = S5 × shares, so the output sums exactly (to floating point) to S5 for every
   LSOA × sex × age (ROL-05).
"""

from __future__ import annotations

import json

import numpy as np
import pandas as pd

from lsc_pop.base import DIMS, ETH, SEXES, load_base
from lsc_pop.config import Config
from lsc_pop.download import raw_file
from lsc_pop.mappings import age_to_code, load_census_age_classifications, load_ethnicity_mapping
from lsc_pop.provenance import Cube, RunContext, logged_step, write_cube
from lsc_pop.validate import check

STAGE = "rollforward"
FALLBACK_LEVELS = {
    0: "direct",
    1: "lsoa_sex_band",
    2: "lsoa_sex_all_ages",
    3: "lsoa_all",
    4: "ltla_sex_age",
}
S5B_BANDS = {
    "0 to 15": (0, 15),
    "16 to 29": (16, 29),
    "30 to 44": (30, 44),
    "45 to 64": (45, 64),
    "65 and over": (65, 90),
}


def _processed(cfg: Config, variant: str | None = None):
    v = variant or cfg.variant
    return cfg.resolve(cfg.paths.processed) / f"mid{cfg.reference_year}_{v}"


# --------------------------------------------------------------------------------------------
# S5
# --------------------------------------------------------------------------------------------


def _read_sheet(path, sheet: str, header: int) -> pd.DataFrame:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise KeyError(f"sheet {sheet!r} not in {path.name}: {wb.sheetnames} (config.mye)")
    rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    df = pd.DataFrame(rows[header + 1 :], columns=rows[header])
    return df.dropna(how="all")


def load_mye(ctx: RunContext, lsoas: list[str]) -> np.ndarray:
    """S5 as an array (L, 2, A) of persons, in ``lsoas`` order."""
    cfg = ctx.cfg
    n_age = cfg.age.max_age + 1
    with logged_step(ctx, "rollforward.load_mye", params=cfg.mye.model_dump()) as step:
        path, entry = raw_file(cfg, cfg.mye.source, cfg.mye.file)
        df = _read_sheet(path, cfg.mye_sheet, cfg.mye.header_row)
        step.input_file(
            f"{cfg.mye.source}/{cfg.mye.file}!{cfg.mye_sheet}", path, entry["sha256"], rows=len(df)
        )
        df = df[df["LSOA 2021 Code"].astype(str).str.startswith("E")]
        cols = [f"{s}{a}" for s in SEXES for a in range(n_age)]
        missing = set(cols) - set(df.columns)
        if missing:
            raise KeyError(f"S5 sheet lacks columns {sorted(missing)[:5]}...; check config.mye")
        df = df.set_index("LSOA 2021 Code")
        vals = df[cols].astype(float)
        got = pd.Index(df.index)
        check(
            ctx,
            "ROL-01",
            "Mid-year estimates cover exactly the lookup LSOAs",
            got.sort_values().equals(pd.Index(lsoas).sort_values()) and got.is_unique,
            stage=STAGE,
            metrics={
                "rows": len(df),
                "missing": len(pd.Index(lsoas).difference(got)),
                "extra": len(got.difference(pd.Index(lsoas))),
            },
        )
        diff = (df["Total"].astype(float) - vals.sum(axis=1)).abs()
        check(
            ctx,
            "ROL-02",
            "Mid-year Total = sum of sex x single-year cells",
            bool((diff == 0).all()),
            stage=STAGE,
            metrics={"lsoas_nonzero": int((diff > 0).sum())},
        )
        arr = vals.loc[lsoas].to_numpy().reshape(len(lsoas), 2, n_age)
        check(
            ctx,
            "ROL-03",
            "Mid-year cells are non-negative integers",
            bool((arr >= 0).all() and (arr == np.round(arr)).all()),
            stage=STAGE,
        )
        step.note(f"{cfg.mye_sheet}: England total {arr.sum():,.0f}")
        long = pd.DataFrame(
            {
                "lsoa21_code": np.repeat(lsoas, 2 * n_age),
                "sex": np.tile(np.repeat(SEXES, n_age), len(lsoas)),
                "age": np.tile(np.arange(n_age), 2 * len(lsoas)),
                "population": arr.ravel(),
            }
        )
        step.output("mye", long)
    return arr


def check_broad_age(ctx: RunContext, mye: np.ndarray, lsoas: list[str]) -> None:
    """S5 aggregated to S5b's broad bands must equal the accredited S5b figures (soft)."""
    cfg = ctx.cfg
    try:
        path, entry = raw_file(cfg, "S5b", "sapelsoabroadage20222024.xlsx")
    except FileNotFoundError:
        check(
            ctx,
            "ROL-04",
            "S5 vs accredited broad-age S5b",
            True,
            stage=STAGE,
            hard=False,
            details="S5b not downloaded; skipped",
        )
        return
    with logged_step(ctx, "rollforward.check_broad_age") as step:
        df = _read_sheet(path, cfg.mye_sheet, cfg.mye.header_row)
        step.input_file(f"S5b!{cfg.mye_sheet}", path, entry["sha256"], rows=len(df))
        df = df[df["LSOA 2021 Code"].astype(str).str.startswith("E")].set_index("LSOA 2021 Code")
        df = df.loc[lsoas]
        worst = 0.0
        for si, s in enumerate(SEXES):
            for label, (lo, hi) in S5B_BANDS.items():
                ours = mye[:, si, lo : hi + 1].sum(1)
                worst = max(worst, float(np.abs(ours - df[f"{s}{label}"].astype(float)).max()))
        check(
            ctx,
            "ROL-04",
            "S5 single-year file agrees with accredited broad-age S5b",
            worst == 0,
            stage=STAGE,
            hard=False,
            metrics={"max_abs_diff": worst},
        )


# --------------------------------------------------------------------------------------------
# Shares
# --------------------------------------------------------------------------------------------


def source_ages(
    variant: str, shift: int, max_age: int, newborn_proxy: list[int]
) -> list[list[int]]:
    """For each target age, the 2021 single years whose ethnic mix it inherits."""
    if variant == "static":
        return [[a] for a in range(max_age + 1)]
    if variant != "cohort":
        raise ValueError(variant)
    out = []
    for a in range(max_age + 1):
        if a < shift:
            out.append(sorted(newborn_proxy))
        elif a < max_age:
            out.append([a - shift])
        else:
            out.append(list(range(max_age - shift, max_age + 1)))
    return out


def compute_shares(
    cfg: Config, base: np.ndarray, ltla_of_lsoa: np.ndarray, src: list[list[int]]
) -> tuple[np.ndarray, np.ndarray]:
    """Shares (L, 2, A, E) summing to 1 over E, and fallback level (L, 2, A) per cell."""
    n_l, _, n_age, _ = base.shape
    a2b = age_to_code(
        load_census_age_classifications(
            cfg.resolve(cfg.age.census_classifications_file), cfg.age.max_age
        ),
        "rm032_5",
        cfg.age.max_age,
    ).to_numpy()
    pooled = np.stack([base[:, :, s, :].sum(2) for s in src], axis=2)  # (L, 2, A, E)
    level = np.zeros((n_l, 2, n_age), dtype=np.int8)

    # Fallback candidates, each (.., E) broadcastable to (L, 2, A, E)
    band_tot = np.stack([base[:, :, a2b == b, :].sum(2) for b in range(1, 6)], axis=2)  # (L,2,5,E)
    src_band = np.array([a2b[s[0]] - 1 for s in src])  # band of each target's source ages
    fb1 = band_tot[:, :, src_band, :]
    fb2 = np.broadcast_to(base.sum(2, keepdims=True), pooled.shape)
    fb3 = np.broadcast_to(base.sum((1, 2), keepdims=True), pooled.shape)
    codes, inv = np.unique(ltla_of_lsoa, return_inverse=True)
    ltla = np.zeros((len(codes),) + base.shape[1:])
    np.add.at(ltla, inv, base)
    ltla_pooled = np.stack([ltla[:, :, s, :].sum(2) for s in src], axis=2)
    fb4 = ltla_pooled[inv]

    counts = pooled.copy()
    for lev, cand in enumerate((fb1, fb2, fb3, fb4), start=1):
        empty = counts.sum(3) <= 0
        if not empty.any():
            break
        counts[empty] = cand[empty]
        level[empty] = lev
    tot = counts.sum(3, keepdims=True)
    if (tot <= 0).any():
        raise ValueError(f"{int((tot <= 0).sum())} cells have no ethnic information at any level")
    return counts / tot, level


# --------------------------------------------------------------------------------------------
# Stage
# --------------------------------------------------------------------------------------------


def roll_forward(
    ctx,
    base_cube: Cube,
    mye: np.ndarray,
    ltla_of_lsoa,
    variant: str,
    newborn_proxy: list[int] | None = None,
) -> tuple:
    """Returns (estimates, shares, fallback level, fallback summary)."""
    cfg = ctx.cfg
    proxy = newborn_proxy if newborn_proxy is not None else cfg.rollforward.newborn_proxy_ages
    src = source_ages(variant, cfg.shift_years, cfg.age.max_age, proxy)
    params = {"variant": variant, "shift_years": cfg.shift_years, "newborn_proxy_ages": proxy}
    with logged_step(ctx, f"rollforward.apply[{variant}]", params=params) as step:
        step.input_cube("base2021", base_cube)
        shares, level = compute_shares(cfg, base_cube.data, ltla_of_lsoa, src)
        est = mye[:, :, :, None] * shares
        step.output_cube("estimates", Cube(est, DIMS, base_cube.coords))
        fb = {
            FALLBACK_LEVELS[k]: {
                "cells": int((level == k).sum()),
                "persons": float(mye[level == k].sum()),
            }
            for k in FALLBACK_LEVELS
        }
        step.note(f"share fallback levels: {json.dumps(fb)}")
    return est, shares, level, fb


def validate_estimates(ctx, est, shares, mye, fb, variant) -> None:
    gap = float(np.abs(est.sum(3) - mye).max())
    check(
        ctx,
        "ROL-05",
        f"[{variant}] Estimates sum to S5 for every LSOA x sex x age",
        gap <= 1e-9 * max(1.0, float(mye.max())),
        stage=STAGE,
        tolerance=1e-9,
        metrics={"max_abs_gap": gap},
    )
    ok = bool(np.isfinite(shares).all() and (shares >= 0).all() and (shares <= 1 + 1e-12).all())
    s_err = float(np.abs(shares.sum(3) - 1).max())
    check(
        ctx,
        "ROL-06",
        f"[{variant}] Shares finite, in [0, 1] and sum to 1",
        ok and s_err < 1e-9,
        stage=STAGE,
        metrics={"max_sum_err": s_err},
    )
    check(
        ctx,
        "ROL-07",
        f"[{variant}] Estimates non-negative/finite; total = S5 total",
        bool(np.isfinite(est).all() and (est >= 0).all() and abs(est.sum() - mye.sum()) < 1e-3),
        stage=STAGE,
        metrics={"total": float(est.sum()), "s5_total": float(mye.sum())},
    )
    check(
        ctx,
        "ROL-08",
        f"[{variant}] Share fallback usage (informational)",
        True,
        stage=STAGE,
        hard=False,
        metrics=fb,
    )


def summarise(cfg, est: np.ndarray, lookup: pd.DataFrame, lsoas: list[str]) -> pd.DataFrame:
    """Totals by geography (England, focus ICB) x eth19 x 10-year band x sex."""
    eth = load_ethnicity_mapping(cfg.resolve(cfg.ethnicity.mapping_file)).set_index("code_19")
    focus = lookup.set_index("lsoa21_code").loc[lsoas, "in_focus_icb"].to_numpy()
    band = np.minimum(np.arange(cfg.age.max_age + 1) // 10 * 10, 90)
    rows = []
    for geo, mask in (("England", np.ones(len(lsoas), bool)), ("Focus ICB", focus)):
        sub = est[mask].sum(0)  # (2, A, E)
        for si, s in enumerate(SEXES):
            for b in np.unique(band):
                v = sub[si, band == b, :].sum(0)
                for e in ETH:
                    rows.append(
                        (
                            geo,
                            s,
                            f"{b}-{b + 9}" if b < 90 else "90+",
                            e,
                            eth.loc[e, "code_6"],
                            float(v[e - 1]),
                        )
                    )
    return pd.DataFrame(
        rows, columns=["geography", "sex", "age_band", "eth19", "eth6", "population"]
    )


def sensitivity(ctx, base_cube, mye, ltla_of_lsoa, lookup, lsoas, default_est) -> pd.DataFrame:
    """Compare variants with the default: static shares, newborn proxy 0-4, seed floors (focus)."""
    from lsc_pop import base as stage_c

    cfg = ctx.cfg
    variants = {f"default ({cfg.variant})": default_est}
    other = "static" if cfg.variant == "cohort" else "cohort"
    variants[other] = roll_forward(ctx, base_cube, mye, ltla_of_lsoa, other)[0]
    if cfg.variant == "cohort":
        variants["cohort, newborn proxy 2021 ages 0-4"] = roll_forward(
            ctx, base_cube, mye, ltla_of_lsoa, "cohort", newborn_proxy=[0, 1, 2, 3, 4]
        )[0]
    frames = [summarise(cfg, v, lookup, lsoas).assign(variant=k) for k, v in variants.items()]

    # Seed floor: refit the focus ICB LSOAs only, roll forward, compare at Focus ICB level.
    rd = lambda n: pd.read_parquet(cfg.resolve(cfg.paths.interim) / "census" / f"{n}.parquet")  # noqa: E731
    seed, ltlas, _ = stage_c.build_seed(
        ctx, rd("seed_age91"), rd("seed_age23"), rd("seed_blocked"), lookup
    )
    eth_m, age_m = stage_c.margin_arrays(cfg, rd("margins"), lsoas)
    idx = np.flatnonzero(lookup.set_index("lsoa21_code").loc[lsoas, "in_focus_icb"].to_numpy())
    li = pd.Index(ltlas).get_indexer(lookup.set_index("lsoa21_code").loc[lsoas, "ltla21_code"])
    sub_lsoas = [lsoas[i] for i in idx]
    sub_lookup = lookup[lookup["lsoa21_code"].isin(sub_lsoas)]
    for fl in cfg.ipf.sensitivity_floors:
        if fl == cfg.ipf.seed_floor:
            continue
        b, _ = stage_c.fit_base(
            ctx,
            seed,
            li[idx],
            eth_m[idx],
            age_m[idx],
            fl,
            label=f"rollforward.sensitivity_fit[floor={fl:g}]",
        )
        cube = Cube(b, DIMS, base_cube.coords | {"lsoa21_code": sub_lsoas})
        est = roll_forward(ctx, cube, mye[idx], ltla_of_lsoa[idx], cfg.variant)[0]
        s = summarise(cfg, est, sub_lookup, sub_lsoas)
        frames.append(s[s["geography"] == "Focus ICB"].assign(variant=f"seed floor {fl:g}"))
    out = pd.concat(frames, ignore_index=True)
    key = ["geography", "sex", "age_band", "eth19", "eth6"]
    ref = out[out["variant"] == f"default ({cfg.variant})"].set_index(key)["population"]
    out = out.join(ref.rename("default"), on=key)
    out["diff"] = out["population"] - out["default"]
    out["pct_diff"] = np.where(out["default"] > 0, out["diff"] / out["default"] * 100, np.nan)
    return out


def run(ctx: RunContext) -> dict:
    from lsc_pop.geography import load_lookup

    cfg = ctx.cfg
    lookup = load_lookup(cfg)
    lsoas = lookup["lsoa21_code"].tolist()
    base_cube = load_base(cfg)
    if base_cube.coords["lsoa21_code"] != lsoas:
        raise ValueError("base cube LSOAs differ from the geography lookup; rerun Stage C")
    ltla_of_lsoa = pd.Index(sorted(lookup["ltla21_code"].unique())).get_indexer(
        lookup["ltla21_code"]
    )

    mye = load_mye(ctx, lsoas)
    check_broad_age(ctx, mye, lsoas)
    est, shares, level, fb = roll_forward(ctx, base_cube, mye, ltla_of_lsoa, cfg.variant)
    validate_estimates(ctx, est, shares, mye, fb, cfg.variant)

    sens = sensitivity(ctx, base_cube, mye, ltla_of_lsoa, lookup, lsoas, est)
    sens.to_csv(ctx.out_dir / "sensitivity_variants.csv", index=False)

    out_dir = _processed(cfg)
    cube = Cube(est, DIMS, base_cube.coords)
    with logged_step(ctx, "rollforward.write") as step:
        step.output_cube("estimates", cube)
        write_cube(
            cube,
            out_dir / "estimates",
            ctx,
            sources=[{"id": i} for i in ("S1", "S2", "S3", cfg.mye.source)],
            description=f"Modelled mid-{cfg.reference_year} population, LSOA x sex x "
            f"single year x ethnic group ({cfg.variant} variant)",
        )
        lv = Cube(level.astype(np.int16), DIMS[:3], {k: base_cube.coords[k] for k in DIMS[:3]})
        write_cube(
            lv,
            out_dir / "share_fallback_level",
            ctx,
            value_name="fallback_level",
            description=f"Share fallback level per LSOA x sex x age: {FALLBACK_LEVELS}",
        )
    return {"fallback": fb, "sensitivity": sens, "total": float(est.sum())}
